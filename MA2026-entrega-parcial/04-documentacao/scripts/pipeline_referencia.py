"""Reproduz em Python o mesmo tratamento feito no Excel e no Power Query.
Serve como gabarito: os números daqui têm que bater com os do Excel e do Power Pivot."""
import csv
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl
import pandas as pd


def ler_xlsx(caminho):
    ws = openpyxl.load_workbook(caminho).active
    linhas = list(ws.iter_rows(values_only=True))
    return [dict(zip(linhas[0], linha), Arquivo_Origem=caminho.name) for linha in linhas[1:]]


def ler_csv(caminho):
    with open(caminho, encoding="utf-8") as arquivo:
        return [dict(linha, Arquivo_Origem=caminho.name) for linha in csv.DictReader(arquivo, delimiter=";")]


def para_data(valor):
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = str(valor).strip()
    if texto[4] == "-":
        return date(int(texto[:4]), int(texto[5:7]), int(texto[8:10]))
    return date(int(texto[6:10]), int(texto[3:5]), int(texto[:2]))


def para_numero(valor):
    if valor is None:
        return 0.0
    if isinstance(valor, (int, float)):
        return float(valor)
    texto = str(valor).replace("R$", "").strip()
    if not texto:
        return 0.0
    return float(texto.replace(".", "").replace(",", "."))


def mapa(caminho, aba):
    ws = openpyxl.load_workbook(caminho)[aba]
    return {original: padrao for original, padrao in ws.iter_rows(min_row=2, values_only=True)}


def tratar(pasta):
    pasta = Path(pasta)
    brutas = ler_xlsx(pasta / "vendas_SP.xlsx") + ler_xlsx(pasta / "vendas_RJ.xlsx") + ler_csv(pasta / "vendas_MG.csv")
    apoio = pasta / "tabelas_apoio.xlsx"
    categoria, cidade, pais = mapa(apoio, "DePara_Categoria"), mapa(apoio, "DePara_Cidade"), mapa(apoio, "DePara_Pais")
    vistos, linhas, auditoria = set(), [], {"linhas_brutas": len(brutas), "vazias": 0, "duplicadas": 0, "invalidas": 0}
    for bruta in brutas:
        if not bruta["ID_Venda"]:
            auditoria["vazias"] += 1
            continue
        if bruta["ID_Venda"] in vistos:
            auditoria["duplicadas"] += 1
            continue
        vistos.add(bruta["ID_Venda"])
        qtd, valor = para_numero(bruta["Qtd"]), para_numero(bruta["Valor_Unitario"])
        if qtd <= 0 or valor <= 0:
            auditoria["invalidas"] += 1
            continue
        cidade_evento = cidade[bruta["Cidade_Evento"].strip().lower()]
        cidade_comprador = cidade[bruta["Cidade_Comprador"].strip().lower()]
        pais_comprador = pais[bruta["Pais_Comprador"].strip().lower()]
        if pais_comprador != "Brasil":
            origem = "Turista Internacional"
        elif cidade_comprador == cidade_evento:
            origem = "Local"
        else:
            origem = "Turista Nacional"
        linhas.append({
            "ID_Venda": bruta["ID_Venda"], "Data_Venda": para_data(bruta["Data_Venda"]),
            "Data_Evento": para_data(bruta["Data_Evento"]), "Evento": bruta["Evento"].strip(),
            "Categoria": categoria[bruta["Categoria"].strip().lower()], "Cidade_Evento": cidade_evento,
            "Tipo_Ingresso": bruta["Tipo_Ingresso"].strip(), "Qtd": int(qtd), "Valor_Unitario": valor,
            "Valor_Total": qtd * valor, "Canal": bruta["Canal"].strip(), "Comprador_ID": bruta["Comprador_ID"],
            "Cidade_Comprador": cidade_comprador, "UF_Comprador": str(bruta["UF_Comprador"]).strip().upper(),
            "Pais_Comprador": pais_comprador, "Origem_Publico": origem, "Arquivo_Origem": bruta["Arquivo_Origem"],
        })
    auditoria["validas"] = len(linhas)
    return pd.DataFrame(linhas), auditoria


def kpis(fato):
    total = fato["Valor_Total"].sum()
    turista = fato["Origem_Publico"] != "Local"
    por_categoria = fato.groupby("Categoria").apply(lambda g: pd.Series({
        "Faturamento": g["Valor_Total"].sum(),
        "Ingressos": g["Qtd"].sum(),
        "Ticket_Medio": g["Valor_Total"].sum() / g["Qtd"].sum(),
        "Faturamento_Turista": g.loc[g["Origem_Publico"] != "Local", "Valor_Total"].sum(),
        "Pct_Faturamento_Turista": g.loc[g["Origem_Publico"] != "Local", "Valor_Total"].sum() / g["Valor_Total"].sum(),
        "Pct_Compradores_Turistas": g.loc[g["Origem_Publico"] != "Local", "Comprador_ID"].nunique() / g["Comprador_ID"].nunique(),
        "Participacao": g["Valor_Total"].sum() / total,
    }), include_groups=False)
    gerais = {
        "Faturamento Total": total,
        "Ingressos Vendidos": fato["Qtd"].sum(),
        "Pedidos": fato["ID_Venda"].nunique(),
        "Ticket Médio (por ingresso)": total / fato["Qtd"].sum(),
        "Ticket Médio (por pedido)": total / fato["ID_Venda"].nunique(),
        "% Faturamento Turista": fato.loc[turista, "Valor_Total"].sum() / total,
        "% Compradores Turistas": fato.loc[turista, "Comprador_ID"].nunique() / fato["Comprador_ID"].nunique(),
    }
    return gerais, por_categoria


def dimensoes(fato, pasta):
    apoio = Path(pasta) / "tabelas_apoio.xlsx"
    dim_evento = pd.read_excel(apoio, "Cadastro_Eventos")
    dim_praca = pd.read_excel(apoio, "Cadastro_Pracas")
    dim_comprador = (fato[["Comprador_ID", "Cidade_Comprador", "UF_Comprador", "Pais_Comprador"]]
                     .drop_duplicates("Comprador_ID").sort_values("Comprador_ID").reset_index(drop=True))
    dim_comprador["Tipo_Comprador"] = dim_comprador["Pais_Comprador"].map(lambda p: "Nacional" if p == "Brasil" else "Estrangeiro")
    inicio = date(min(fato["Data_Venda"].min(), fato["Data_Evento"].min()).year, 1, 1)
    fim = date(max(fato["Data_Venda"].max(), fato["Data_Evento"].max()).year, 12, 31)
    nomes_mes = ["jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago", "set", "out", "nov", "dez"]
    nomes_dia = ["segunda", "terça", "quarta", "quinta", "sexta", "sábado", "domingo"]
    dias = [inicio + timedelta(days=n) for n in range((fim - inicio).days + 1)]
    dim_calendario = pd.DataFrame({
        "Data": dias, "Ano": [d.year for d in dias], "Mes": [d.month for d in dias],
        "Nome_Mes": [nomes_mes[d.month - 1] for d in dias], "Ano_Mes": [f"{d.year}-{d.month:02d}" for d in dias],
        "Trimestre": [f"T{(d.month - 1) // 3 + 1}" for d in dias], "Dia_Semana": [nomes_dia[d.weekday()] for d in dias],
        "Fim_de_Semana": ["Sim" if d.weekday() >= 5 else "Não" for d in dias],
    })
    return dim_evento, dim_praca, dim_comprador, dim_calendario
