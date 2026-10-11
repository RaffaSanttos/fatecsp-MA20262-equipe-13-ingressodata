"""Gera os dados brutos SIMULADOS do IngressoData Brasil (seed fixa = reproduzível).

Saídas (em dados-brutos/):
  vendas_SP.xlsx   datas em texto dd/mm/aaaa, valores em texto "R$ 1.250,00"
  vendas_RJ.xlsx   datas e valores como números do Excel
  vendas_MG.csv    separador ";", datas aaaa-mm-dd, decimal com vírgula
  tabelas_apoio.xlsx  De-Para (categoria, cidade, país) + cadastro de eventos e praças

Os problemas de qualidade (grafias diferentes, duplicados, linhas vazias, quantidade
ou valor faltando) são inseridos de propósito para o tratamento no Excel e no Power Query.
"""
import csv
import random
import sys
import unicodedata
from datetime import date, timedelta
from pathlib import Path

import openpyxl

SEED = 20260918
random.seed(SEED)

PRACAS = [
    # cidade, UF, regiao, lat, lon, arquivo
    ("São Paulo", "SP", "Sudeste", -23.5505, -46.6333, "SP"),
    ("Rio de Janeiro", "RJ", "Sudeste", -22.9068, -43.1729, "RJ"),
    ("Belo Horizonte", "MG", "Sudeste", -19.9167, -43.9345, "MG"),
]

# evento, categoria, cidade, local, ao_ar_livre, data_evento (None = sessões diárias)
EVENTOS = [
    ("Festival Sonora SP 2025", "Festival", "São Paulo", "Autódromo de Interlagos", "Sim", date(2025, 11, 15)),
    ("Festival Sonora SP 2026", "Festival", "São Paulo", "Autódromo de Interlagos", "Sim", date(2026, 4, 18)),
    ("Rio Beat Festival 2026", "Festival", "Rio de Janeiro", "Parque Olímpico", "Sim", date(2026, 1, 24)),
    ("Festival Montanhas BH 2025", "Festival", "Belo Horizonte", "Esplanada do Mineirão", "Sim", date(2025, 8, 23)),
    ("Turnê Internacional Pop - SP", "Show", "São Paulo", "Estádio Allianz Parque", "Sim", date(2025, 10, 4)),
    ("Show Nacional MPB - Rio", "Show", "Rio de Janeiro", "Arena Jeunesse", "Não", date(2025, 6, 14)),
    ("Turnê Internacional Rock - Rio", "Show", "Rio de Janeiro", "Estádio Nilton Santos", "Sim", date(2026, 5, 9)),
    ("Show Sertanejo - BH", "Show", "Belo Horizonte", "Mineirão", "Sim", date(2026, 2, 28)),
    ("Cinema Paulista - Sessões", "Cinema", "São Paulo", "Cine Shopping Paulista", "Não", None),
    ("Cinema Barra - Sessões", "Cinema", "Rio de Janeiro", "Cine Barra Shopping", "Não", None),
    ("Cinema Savassi - Sessões", "Cinema", "Belo Horizonte", "Cine Savassi", "Não", None),
]

CIDADES_BR = [
    ("São Paulo", "SP", "01"), ("Campinas", "SP", "13"), ("Santos", "SP", "11"),
    ("Rio de Janeiro", "RJ", "20"), ("Niterói", "RJ", "24"), ("Belo Horizonte", "MG", "30"),
    ("Contagem", "MG", "32"), ("Curitiba", "PR", "80"), ("Porto Alegre", "RS", "90"),
    ("Salvador", "BA", "40"), ("Brasília", "DF", "70"), ("Recife", "PE", "50"),
    ("Goiânia", "GO", "74"), ("Florianópolis", "SC", "88"), ("Vitória", "ES", "29"),
]
CIDADES_EXT = [("Buenos Aires", "Argentina"), ("Santiago", "Chile"), ("Miami", "EUA"), ("Lisboa", "Portugal")]
VIZINHAS = {"São Paulo": ["Campinas", "Santos"], "Rio de Janeiro": ["Niterói"], "Belo Horizonte": ["Contagem"]}

# (participação local, participação estrangeiro) por categoria
PERFIL_PUBLICO = {"Festival": (0.55, 0.08), "Show": (0.68, 0.05), "Cinema": (0.95, 0.005)}
PRECO_BASE = {"Festival": (650, 1150), "Show": (220, 780), "Cinema": (28, 48)}
QTD_VENDAS = {"Festival": 230, "Show": 150, "Cinema": 170}

VARIANTES_CATEGORIA = {
    "Show": ["Show", "SHOW", "show", "Show ", "Shows"],
    "Festival": ["Festival", "FESTIVAL", "festival", "Festivais", "Fest."],
    "Cinema": ["Cinema", "CINEMA", "cinema", "Cine", "Sessão Cinema"],
}
VARIANTES_PRACA = {
    "São Paulo": ["São Paulo", "Sao Paulo", "SÃO PAULO", "SP - São Paulo", "São Paulo "],
    "Rio de Janeiro": ["Rio de Janeiro", "Rio De Janeiro", "RIO DE JANEIRO", "Rio", "rio de janeiro "],
    "Belo Horizonte": ["Belo Horizonte", "BELO HORIZONTE", "Belo Horizonte ", "BH", "belo horizonte"],
}
VARIANTES_PAIS = {"Brasil": ["Brasil", "brasil", "BR", "BRASIL"], "Argentina": ["Argentina", "ARGENTINA"],
                  "Chile": ["Chile", "chile"], "EUA": ["EUA", "Estados Unidos", "USA"], "Portugal": ["Portugal", "PORTUGAL"]}

PERIODO_INICIO, PERIODO_FIM = date(2025, 1, 1), date(2026, 6, 30)
CABECALHOS = ["ID_Venda", "Data_Venda", "Data_Evento", "Evento", "Categoria", "Cidade_Evento", "Tipo_Ingresso",
              "Qtd", "Valor_Unitario", "Canal", "Comprador_ID", "Cidade_Comprador", "UF_Comprador",
              "Pais_Comprador", "CEP_Comprador"]


def sem_acento(texto):
    return "".join(c for c in unicodedata.normalize("NFD", texto) if unicodedata.category(c) != "Mn")


def variantes_cidade(cidade):
    return [cidade, cidade.upper(), cidade.lower(), sem_acento(cidade), cidade + " "]


def escolher(variantes, chance_limpo=0.55):
    return variantes[0] if random.random() < chance_limpo else random.choice(variantes)


def criar_compradores(qtd=2500):
    compradores = []
    for numero in range(1, qtd + 1):
        if random.random() < 0.06:
            cidade, pais = random.choice(CIDADES_EXT)
            compradores.append((f"C{numero:05d}", cidade, "EX", pais, ""))
        else:
            cidade, uf, prefixo = random.choice(CIDADES_BR)
            cep = f"{prefixo}{random.randint(0, 999):03d}-{random.randint(0, 999):03d}"
            compradores.append((f"C{numero:05d}", cidade, uf, "Brasil", cep))
    por_cidade = {}
    for comprador in compradores:
        por_cidade.setdefault(comprador[1], []).append(comprador)
    return compradores, por_cidade


def sortear_comprador(categoria, praca, por_cidade):
    p_local, p_estrangeiro = PERFIL_PUBLICO[categoria]
    sorteio = random.random()
    if sorteio < p_local:
        cidades = [praca]
    elif sorteio < p_local + p_estrangeiro:
        cidades = [c for c, _ in CIDADES_EXT]
    else:
        cidades = [c for c, _, _ in CIDADES_BR if c != praca]
        cidades += VIZINHAS[praca] * 3  # turistas de cidades vizinhas são mais comuns
    return random.choice(por_cidade[random.choice(cidades)])


def data_da_venda(evento_data):
    if evento_data is None:
        dias = (PERIODO_FIM - PERIODO_INICIO).days
        venda = PERIODO_INICIO + timedelta(days=random.randint(0, dias))
        return venda, venda + timedelta(days=random.randint(0, 3))
    abertura = 120
    # picos: abertura das vendas e última semana
    sorteio = random.random()
    if sorteio < 0.35:
        antecedencia = random.randint(100, abertura)
    elif sorteio < 0.65:
        antecedencia = random.randint(0, 7)
    else:
        antecedencia = random.randint(8, 99)
    return evento_data - timedelta(days=antecedencia), evento_data


def gerar_vendas():
    _, por_cidade = criar_compradores()
    vendas = {"SP": [], "RJ": [], "MG": []}
    sigla = {p[0]: p[5] for p in PRACAS}
    contador = {"SP": 0, "RJ": 0, "MG": 0}
    for evento, categoria, praca, _, _, data_evento in EVENTOS:
        arquivo = sigla[praca]
        for _ in range(QTD_VENDAS[categoria]):
            contador[arquivo] += 1
            data_venda, data_ev = data_da_venda(data_evento)
            comprador = sortear_comprador(categoria, praca, por_cidade)
            minimo, maximo = PRECO_BASE[categoria]
            tipo = random.choices(["Inteira", "Meia", "VIP"], [0.55, 0.35, 0.10])[0]
            if categoria == "Cinema" and tipo == "VIP":
                tipo = "Inteira"
            preco = round(random.uniform(minimo, maximo) / 5) * 5
            preco = preco * 2 if tipo == "VIP" else preco / 2 if tipo == "Meia" else preco
            qtd = random.choices([1, 2, 3, 4], [0.45, 0.35, 0.12, 0.08])[0]
            if comprador[2] != "EX" and comprador[1] != praca:
                qtd = random.choices([1, 2, 3, 4], [0.30, 0.45, 0.15, 0.10])[0]
            vendas[arquivo].append({
                "ID_Venda": f"{arquivo}-{contador[arquivo]:06d}",
                "Data_Venda": data_venda, "Data_Evento": data_ev, "Evento": evento,
                "Categoria": escolher(VARIANTES_CATEGORIA[categoria]),
                "Cidade_Evento": escolher(VARIANTES_PRACA[praca]),
                "Tipo_Ingresso": tipo, "Qtd": qtd, "Valor_Unitario": float(preco),
                "Canal": random.choices(["Site", "App", "Bilheteria", "Parceiro"], [0.45, 0.35, 0.12, 0.08])[0],
                "Comprador_ID": comprador[0], "Cidade_Comprador": escolher(variantes_cidade(comprador[1]), 0.6),
                "UF_Comprador": comprador[2] if random.random() < 0.9 else comprador[2].lower(),
                "Pais_Comprador": escolher(VARIANTES_PAIS[comprador[3]], 0.7), "CEP_Comprador": comprador[4],
            })
    for arquivo, linhas in vendas.items():
        linhas.sort(key=lambda linha: linha["Data_Venda"])
        for linha in random.sample(linhas, int(len(linhas) * 0.01)):
            linha["Qtd"] = None if random.random() < 0.5 else 0
        for linha in random.sample(linhas, int(len(linhas) * 0.006)):
            linha["Valor_Unitario"] = None
        duplicadas = random.sample(linhas, int(len(linhas) * 0.025))
        for linha in duplicadas:  # cópia exata inserida depois da original
            posicao = linhas.index(linha) + random.randint(1, 40)
            linhas.insert(min(posicao, len(linhas)), dict(linha))
        for _ in range(4):  # linhas totalmente vazias
            linhas.insert(random.randint(1, len(linhas) - 1), None)
    return vendas


def brl(valor):
    return "R$ " + f"{valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def salvar_xlsx(caminho, linhas, formato):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Vendas"
    ws.append(CABECALHOS)
    for linha in linhas:
        if linha is None:
            ws.append([None] * len(CABECALHOS))
            continue
        valores = []
        for coluna in CABECALHOS:
            valor = linha[coluna]
            if formato == "SP" and coluna in ("Data_Venda", "Data_Evento"):
                valor = valor.strftime("%d/%m/%Y")
            elif formato == "SP" and coluna == "Valor_Unitario" and valor is not None:
                valor = brl(valor)
            valores.append(valor)
        ws.append(valores)
        if formato == "RJ":
            ws.cell(ws.max_row, 2).number_format = "dd/mm/yyyy"
            ws.cell(ws.max_row, 3).number_format = "dd/mm/yyyy"
    wb.save(caminho)


def salvar_csv(caminho, linhas):
    with open(caminho, "w", newline="", encoding="utf-8") as arquivo:
        escritor = csv.writer(arquivo, delimiter=";")
        escritor.writerow(CABECALHOS)
        for linha in linhas:
            if linha is None:
                escritor.writerow([""] * len(CABECALHOS))
                continue
            valores = []
            for coluna in CABECALHOS:
                valor = linha[coluna]
                if coluna in ("Data_Venda", "Data_Evento"):
                    valor = valor.isoformat()
                elif coluna == "Valor_Unitario" and valor is not None:
                    valor = f"{valor:.2f}".replace(".", ",")
                valores.append("" if valor is None else valor)
            escritor.writerow(valores)


def tabelas_de_para():
    categoria = sorted({(v.strip().lower(), padrao) for padrao, vs in VARIANTES_CATEGORIA.items() for v in vs})
    cidades = [c for c, _, _ in CIDADES_BR] + [c for c, _ in CIDADES_EXT]
    cidade = {(v.strip().lower(), c) for c in cidades for v in variantes_cidade(c)}
    cidade |= {(v.strip().lower(), padrao) for padrao, vs in VARIANTES_PRACA.items() for v in vs}
    pais = sorted({(v.strip().lower(), padrao) for padrao, vs in VARIANTES_PAIS.items() for v in vs})
    return categoria, sorted(cidade), pais


def salvar_tabelas_apoio(caminho):
    categoria, cidade, pais = tabelas_de_para()
    wb = openpyxl.Workbook()
    abas = [("DePara_Categoria", ["Original", "Padrao"], categoria),
            ("DePara_Cidade", ["Original", "Padrao"], cidade),
            ("DePara_Pais", ["Original", "Padrao"], pais),
            ("Cadastro_Eventos", ["Evento", "Categoria", "Cidade_Evento", "Local", "Ao_Ar_Livre", "Data_Evento"],
             [(e, c, p, l, a, d) for e, c, p, l, a, d in EVENTOS]),
            ("Cadastro_Pracas", ["Cidade", "UF", "Regiao", "Latitude", "Longitude"],
             [p[:5] for p in PRACAS])]
    wb.remove(wb.active)
    for nome, cabecalho, linhas in abas:
        ws = wb.create_sheet(nome)
        ws.append(cabecalho)
        for linha in linhas:
            ws.append(list(linha))
        if nome == "Cadastro_Eventos":
            for celula in ws["F"][1:]:
                celula.number_format = "dd/mm/yyyy"
    wb.save(caminho)


def main(pasta_saida):
    pasta = Path(pasta_saida)
    pasta.mkdir(parents=True, exist_ok=True)
    vendas = gerar_vendas()
    salvar_xlsx(pasta / "vendas_SP.xlsx", vendas["SP"], "SP")
    salvar_xlsx(pasta / "vendas_RJ.xlsx", vendas["RJ"], "RJ")
    salvar_csv(pasta / "vendas_MG.csv", vendas["MG"])
    salvar_tabelas_apoio(pasta / "tabelas_apoio.xlsx")
    return vendas


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "dados-brutos")
