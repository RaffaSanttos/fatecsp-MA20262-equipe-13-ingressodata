"""Gera 03-powerpivot (tabelas do Star Schema + gabarito) e a carga SQL a partir do pipeline de referência."""
import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

import pipeline_referencia as ref

FONTE = "Arial"
CABECALHO = PatternFill("solid", fgColor="1F3A5F")
FORMATOS = {"Data": "dd/mm/yyyy", "Data_Venda": "dd/mm/yyyy", "Data_Evento": "dd/mm/yyyy",
            "Valor_Unitario": '"R$" #,##0.00', "Valor_Total": '"R$" #,##0.00'}

MEDIDAS = [
    ("Faturamento Total", "SUM ( Fato_Vendas_Ingressos[Valor_Total] )", "Faturamento Total"),
    ("Ingressos Vendidos", "SUM ( Fato_Vendas_Ingressos[Qtd] )", "Ingressos Vendidos"),
    ("Pedidos", "DISTINCTCOUNT ( Fato_Vendas_Ingressos[ID_Venda] )", "Pedidos"),
    ("Ticket Médio", "DIVIDE ( [Faturamento Total], [Ingressos Vendidos] )", "Ticket Médio (por ingresso)"),
    ("Ticket Médio por Pedido", "DIVIDE ( [Faturamento Total], [Pedidos] )", "Ticket Médio (por pedido)"),
    ("Faturamento Turista",
     'CALCULATE ( [Faturamento Total], Fato_Vendas_Ingressos[Origem_Publico] <> "Local" )', None),
    ("% Faturamento Turista", "DIVIDE ( [Faturamento Turista], [Faturamento Total] )", "% Faturamento Turista"),
    ("Compradores", "DISTINCTCOUNT ( Fato_Vendas_Ingressos[Comprador_ID] )", None),
    ("Compradores Turistas",
     'CALCULATE ( [Compradores], Fato_Vendas_Ingressos[Origem_Publico] <> "Local" )', None),
    ("% Compradores Turistas", "DIVIDE ( [Compradores Turistas], [Compradores] )", "% Compradores Turistas"),
    ("Participação Categoria",
     "DIVIDE ( [Faturamento Total], CALCULATE ( [Faturamento Total], ALL ( Fato_Vendas_Ingressos[Categoria] ), "
     "ALL ( Dim_Evento[Categoria] ) ) )", None),
    ("Faturamento Ao Ar Livre", 'CALCULATE ( [Faturamento Total], Dim_Evento[Ao_Ar_Livre] = "Sim" )', None),
    ("Faturamento Mês Anterior",
     "CALCULATE ( [Faturamento Total], DATEADD ( Dim_Calendario[Data], -1, MONTH ) )", None),
    ("Variação Mensal %",
     "DIVIDE ( [Faturamento Total] - [Faturamento Mês Anterior], [Faturamento Mês Anterior] )", None),
]

RELACIONAMENTOS = [
    ("Fato_Vendas_Ingressos", "Evento", "Dim_Evento", "Evento"),
    ("Fato_Vendas_Ingressos", "Cidade_Evento", "Dim_Praca", "Cidade"),
    ("Fato_Vendas_Ingressos", "Comprador_ID", "Dim_Comprador", "Comprador_ID"),
    ("Fato_Vendas_Ingressos", "Data_Venda", "Dim_Calendario", "Data"),
]


def aba_tabela(wb, nome, df):
    ws = wb.create_sheet(nome)
    ws.append(list(df.columns))
    for linha in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and v != v) else v for v in linha])
    for indice, coluna in enumerate(df.columns, start=1):
        largura = max(12, min(32, len(coluna) + 4))
        ws.column_dimensions[get_column_letter(indice)].width = largura
        if coluna in FORMATOS:
            for celula in ws[get_column_letter(indice)][1:]:
                celula.number_format = FORMATOS[coluna]
    for celula in ws[1]:
        celula.font = Font(name=FONTE, bold=True, color="FFFFFF")
        celula.fill = CABECALHO
    tabela = Table(displayName=nome, ref=f"A1:{get_column_letter(len(df.columns))}{len(df) + 1}")
    tabela.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(tabela)
    ws.freeze_panes = "A2"


def aba_modelo(wb, gerais):
    ws = wb.active
    ws.title = "Modelo_e_DAX"
    titulo = Font(name=FONTE, bold=True, size=13, color="1F3A5F")
    ws.append(["IngressoData Brasil — Etapa 2: Power Pivot (Star Schema) e DAX"])
    ws["A1"].font = Font(name=FONTE, bold=True, size=15, color="1F3A5F")
    ws.append(["Cada aba é uma Tabela do Excel. Power Pivot > Adicionar ao Modelo de Dados, criar os relacionamentos "
               "e colar as medidas abaixo. Caminho preferido: carregar as consultas do 02-powerquery direto no Modelo."])
    ws.append([])
    ws.append(["Relacionamentos (muitos → um)"])
    ws.cell(ws.max_row, 1).font = titulo
    ws.append(["Tabela (lado muitos)", "Coluna", "Tabela (lado um)", "Coluna"])
    for relacao in RELACIONAMENTOS:
        ws.append(list(relacao))
    ws.append([])
    ws.append(["Medidas DAX (tabela Fato_Vendas_Ingressos)"])
    ws.cell(ws.max_row, 1).font = titulo
    ws.append(["Medida", "Fórmula DAX", "Valor esperado (total geral)"])
    linha_cabecalho_medidas = ws.max_row
    for nome, formula, chave in MEDIDAS:
        esperado = gerais.get(chave) if chave else None
        ws.append([nome, f"{nome} := {formula}", esperado])
        if chave and chave.startswith("%"):
            ws.cell(ws.max_row, 3).number_format = "0.00%"
        elif chave and ("Faturamento" in chave or "Ticket" in chave):
            ws.cell(ws.max_row, 3).number_format = '"R$" #,##0.00'
    for linha in (5, linha_cabecalho_medidas):
        for celula in ws[linha]:
            celula.font = Font(name=FONTE, bold=True, color="FFFFFF")
            celula.fill = CABECALHO
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 110
    ws.column_dimensions["C"].width = 26
    ws.column_dimensions["D"].width = 16
    for linha in ws.iter_rows():
        for celula in linha:
            if celula.font.name != FONTE:
                celula.font = Font(name=FONTE, bold=celula.font.bold, color=celula.font.color, size=celula.font.size)
            celula.alignment = Alignment(wrap_text=celula.column == 2 and celula.row > 2, vertical="top")


def sql_texto(valor):
    if valor is None or (isinstance(valor, float) and valor != valor):
        return "NULL"
    if isinstance(valor, (int, float)):
        return repr(round(valor, 6)) if isinstance(valor, float) else str(valor)
    if hasattr(valor, "isoformat"):
        return f"'{valor.isoformat()[:10]}'"
    return "'" + str(valor).replace("'", "''") + "'"


def gerar_carga_sql(caminho, fato, dim_evento, dim_praca, dim_comprador):
    blocos = ["-- Carga gerada a partir da base tratada (mesmo resultado do Power Query). SQLite e PostgreSQL.",
              "BEGIN;"]
    tabelas = [
        ("dim_praca", ["cidade", "uf", "regiao", "latitude", "longitude"], dim_praca.values.tolist()),
        ("dim_evento", ["evento", "categoria", "cidade_evento", "local_evento", "ao_ar_livre", "data_evento"],
         [[e, c, p, l, a, None if d != d else d] for e, c, p, l, a, d in dim_evento.values.tolist()]),
        ("dim_comprador", ["comprador_id", "cidade", "uf", "pais", "tipo_comprador"], dim_comprador.values.tolist()),
        ("fato_vendas_ingressos",
         ["id_venda", "data_venda", "data_evento", "evento", "categoria", "cidade_evento", "tipo_ingresso", "qtd",
          "valor_unitario", "valor_total", "canal", "comprador_id", "origem_publico", "arquivo_origem"],
         fato[["ID_Venda", "Data_Venda", "Data_Evento", "Evento", "Categoria", "Cidade_Evento", "Tipo_Ingresso",
               "Qtd", "Valor_Unitario", "Valor_Total", "Canal", "Comprador_ID", "Origem_Publico",
               "Arquivo_Origem"]].values.tolist()),
    ]
    for nome, colunas, linhas in tabelas:
        blocos.append(f"\n-- {nome}: {len(linhas)} linhas")
        for inicio in range(0, len(linhas), 200):
            valores = ",\n".join("(" + ", ".join(sql_texto(v) for v in linha) + ")" for linha in linhas[inicio:inicio + 200])
            blocos.append(f"INSERT INTO {nome} ({', '.join(colunas)}) VALUES\n{valores};")
    blocos.append("COMMIT;")
    Path(caminho).write_text("\n".join(blocos) + "\n", encoding="utf-8")


def main(pasta_brutos, pasta_saida):
    saida = Path(pasta_saida)
    saida.mkdir(parents=True, exist_ok=True)
    fato, auditoria = ref.tratar(pasta_brutos)
    gerais, por_categoria = ref.kpis(fato)
    dim_evento, dim_praca, dim_comprador, dim_calendario = ref.dimensoes(fato, pasta_brutos)

    wb = openpyxl.Workbook()
    aba_modelo(wb, {k: float(v) for k, v in gerais.items()})
    aba_tabela(wb, "Fato_Vendas_Ingressos", fato)
    aba_tabela(wb, "Dim_Evento", dim_evento)
    aba_tabela(wb, "Dim_Praca", dim_praca)
    aba_tabela(wb, "Dim_Comprador", dim_comprador)
    aba_tabela(wb, "Dim_Calendario", dim_calendario)
    wb.save(saida / "IngressoData_03_PowerPivot.xlsx")

    gerar_carga_sql(saida / "02_carga_dados.sql", fato, dim_evento, dim_praca, dim_comprador)
    gabarito = {"auditoria": auditoria, "gerais": {k: round(float(v), 4) for k, v in gerais.items()},
                "por_categoria": json.loads(por_categoria.round(4).to_json(orient="index")),
                "por_praca_origem": {f"{c}|{o}": float(v) for (c, o), v in
                                     fato.groupby(["Cidade_Evento", "Origem_Publico"])["Valor_Total"].sum().items()},
                "linhas_dim": {"Dim_Evento": len(dim_evento), "Dim_Praca": len(dim_praca),
                               "Dim_Comprador": len(dim_comprador), "Dim_Calendario": len(dim_calendario)}}
    (saida / "gabarito.json").write_text(json.dumps(gabarito, ensure_ascii=False, indent=2), encoding="utf-8")
    medidas = "\n\n".join(f"{nome} :=\n{formula}" for nome, formula, _ in MEDIDAS)
    (saida / "medidas.dax").write_text(
        "// Medidas DAX do IngressoData Brasil — colar no Power Pivot (Medidas > Nova Medida) ou no Power BI.\n"
        "// Tabela de destino: Fato_Vendas_Ingressos. Dim_Calendario deve ser marcada como Tabela de Data.\n\n"
        + medidas + "\n", encoding="utf-8")
    return gabarito


if __name__ == "__main__":
    print(json.dumps(main(sys.argv[1], sys.argv[2])["gerais"], ensure_ascii=False, indent=2))
